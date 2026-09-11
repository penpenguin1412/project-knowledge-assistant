"""Run: python tests.py. Isolated local DB, no network/model calls or file deletion."""
import io
import os
import unittest
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

os.environ["ASSISTANT_DB"] = str(Path(__file__).parent / "data" / f"test-{uuid.uuid4().hex}.db")
os.environ["LLM_ENABLED"] = "0"
os.environ.pop("ASSISTANT_TOKEN", None)

from fastapi.testclient import TestClient
from pypdf import PdfWriter
from docx import Document
import app as server
import httpx
from core import MAX_BYTES, parse_file
from model import model_status
from model import generate


class FlowTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(server.app, headers={"X-Assistant-Request": "1"})
        self.project = self.client.post('/api/projects', json={"name": "测试项目"}).json()["id"]
        self.path = f'/api/projects/{self.project}'

    def tearDown(self):
        self.client.close()

    def upload(self, text="项目交付物：现状问题地图。", name="sample.md"):
        return self.client.post(self.path+'/documents', files={"file": (name, text.encode(), "text/plain")})

    def draft(self):
        return self.client.post(self.path+'/drafts', json={"text": "待办：完成地图 | 负责人：许宁 | 截止：2026-09-18", "mode": "rules"}).json()

    def test_evidence_and_original(self):
        doc = self.upload().json()
        r = self.client.post(self.path+'/ask', json={"question": "项目交付物有哪些？"}).json()
        self.assertFalse(r['abstained'])
        self.assertEqual(r['mode'], 'evidence')
        self.assertIn('不是大模型', r['answer'])
        self.assertIn('现状问题地图', r['citations'][0]['quote'])
        original = self.client.get(self.path+f"/documents/{doc['id']}").json()
        self.assertEqual(original['chunks'][0]['text'], r['citations'][0]['quote'])
        self.assertEqual(len(self.client.get(self.path+'/runs').json()), 1)

    def test_project_isolation(self):
        doc = self.upload().json()
        other = self.client.post('/api/projects', json={"name":"另一个项目"}).json()['id']
        self.assertEqual(self.client.get(f'/api/projects/{other}/documents').json(), [])
        self.assertEqual(self.client.get(f'/api/projects/{other}/documents/{doc["id"]}').status_code,404)
        r=self.client.post(f'/api/projects/{other}/ask',json={"question":"项目交付物"}).json()
        self.assertTrue(r['abstained'])

    def test_archive_and_restore(self):
        doc=self.upload().json()
        url=self.path+f'/documents/{doc["id"]}'
        self.client.patch(url,json={"archived":True})
        self.assertTrue(self.client.post(self.path+'/ask',json={"question":"项目交付物"}).json()['abstained'])
        self.client.patch(url,json={"archived":False})
        self.assertFalse(self.client.post(self.path+'/ask',json={"question":"项目交付物"}).json()['abstained'])

    def test_duplicate_upload(self):
        self.assertEqual(self.upload().status_code,201)
        self.assertEqual(self.upload(name="renamed.txt").status_code,409)

    def test_upload_boundaries(self):
        for name, raw in [('bad.exe',b'abcd'),('empty.txt',b''),('bad.txt',b'\xff'),('bad.pdf',b'broken'),('null.md',b'x\x00x'),('big.md',b'a'*(MAX_BYTES+1))]:
            with self.subTest(name=name):
                result=self.client.post(self.path+'/documents',files={'file':(name,raw)})
                self.assertIn(result.status_code,[413,422])
        self.assertEqual(self.client.get(self.path+'/documents').json(),[])

    def test_pdf_and_docx(self):
        from pypdf.generic import DecodedStreamObject, NameObject, DictionaryObject
        writer=PdfWriter();page=writer.add_blank_page(width=300,height=300)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
        stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 10 250 Td (Project deadline: 2026-09-30) Tj ET')
        page[NameObject('/Contents')]=stream
        raw=io.BytesIO();writer.write(raw)
        chunks=parse_file('a.pdf',raw.getvalue())
        self.assertIn('2026-09-30',chunks[0]['text']);self.assertIn('第 1 页',chunks[0]['location'])
        doc=Document();doc.add_paragraph('项目资料');doc.add_table(rows=1,cols=1).cell(0,0).text='表格数据'
        raw=io.BytesIO();doc.save(raw)
        self.assertTrue(any('表格数据' in c['text'] for c in parse_file('a.docx',raw.getvalue())))
        writer=PdfWriter();writer.add_blank_page(width=100,height=100);raw=io.BytesIO();writer.write(raw)
        with self.assertRaisesRegex(ValueError,'未发现文字'):parse_file('scan.pdf',raw.getvalue())

    def test_request_guards(self):
        client=TestClient(server.app)
        self.assertEqual(client.post('/api/projects',json={'name':'bad'}).status_code,403)
        self.assertEqual(self.client.get('/api/projects',headers={'Origin':'https://evil.example'}).status_code,403)
        self.assertEqual(self.client.get('/api/projects',headers={'Host':'evil.example'}).status_code,400)
        with patch.dict(os.environ,{'ASSISTANT_TOKEN':'test-only-access'}):
            self.assertEqual(self.client.get('/api/projects').status_code,401)
            self.assertEqual(self.client.get('/api/projects',headers={'Authorization':'Bearer test-only-access'}).status_code,200)
        client.close()

    def test_body_stream_limit(self):
        def body():
            for _ in range(7):yield b'x'*(1024*1024)
        result=self.client.post('/api/projects',content=body(),headers={'Content-Type':'application/json'})
        self.assertIn(result.status_code,[400,413])
        self.assertEqual(self.client.get('/api/health').status_code,200)

    def test_input_validation(self):
        self.assertEqual(self.client.post('/api/projects',json={'name':'  '}).status_code,422)
        self.assertEqual(self.client.post(self.path+'/ask',json={'question':'x','unknown':1}).status_code,422)
        self.assertEqual(self.client.get('/api/projects/999999/tasks').status_code,404)

    def test_confirm_before_storage_and_idempotency(self):
        d=self.draft()
        self.assertEqual(self.client.get(self.path+'/tasks').json(),[])
        self.assertEqual(self.client.get(self.path+'/drafts').json()[0]['id'],d['id'])
        url=self.path+f'/drafts/{d["id"]}/confirm'
        self.assertEqual(self.client.post(url,json={'confirmed':False,'items':d['items']}).status_code,422)
        self.assertEqual(self.client.post(url,json={'confirmed':True,'items':d['items']}).status_code,200)
        self.assertEqual(self.client.post(url,json={'confirmed':True,'items':d['items']}).status_code,409)
        self.assertEqual(len(self.client.get(self.path+'/tasks').json()),1)

    def test_concurrent_confirmation(self):
        d=self.draft();url=self.path+f'/drafts/{d["id"]}/confirm'
        with ThreadPoolExecutor(max_workers=2) as pool:
            statuses=list(pool.map(lambda _:self.client.post(url,json={'confirmed':True,'items':d['items']}).status_code,range(2)))
        self.assertEqual(sorted(statuses),[200,409])

    def test_draft_source_and_project_guard(self):
        d=self.draft();items=[{**d['items'][0],'quote':'伪造原文'}]
        self.assertEqual(self.client.post(self.path+f'/drafts/{d["id"]}/confirm',json={'confirmed':True,'items':items}).status_code,422)
        self.assertEqual(self.client.post(f'/api/projects/999/drafts/{d["id"]}/confirm',json={'confirmed':True,'items':d['items']}).status_code,404)
        self.assertEqual(self.client.get(self.path+'/tasks').json(),[])

    def test_unknown_fields_and_edit(self):
        d=self.client.post(self.path+'/drafts',json={'text':'待办：联系社区 | 负责人：待定 | 截止：待定'}).json()
        self.assertIsNone(d['items'][0]['due_date']);self.assertEqual(d['items'][0]['owner'],'')
        self.client.post(self.path+f'/drafts/{d["id"]}/confirm',json={'confirmed':True,'items':d['items']})
        t=self.client.get(self.path+'/tasks').json()[0]
        r=self.client.patch(self.path+f'/tasks/{t["id"]}',json={'title':t['title'],'owner':'人工指定','due_date':'2026-09-25','status':'done'})
        self.assertEqual(r.status_code,200)
        self.assertEqual(self.client.get(self.path+'/tasks').json()[0]['status'],'done')

    def test_rules_do_not_invent(self):
        d=self.client.post(self.path+'/drafts',json={'text':'已完成地图。下周再讨论预算，不采购无人机。'}).json()
        self.assertEqual(d['items'],[])
        self.assertEqual(self.client.post(self.path+'/drafts',json={'text':'待办：地图 | 负责人：甲 | 截止：2026-02-30'}).status_code,422)

    def test_missing_model_is_explicit(self):
        self.upload()
        result=self.client.post(self.path+'/ask',json={'question':'项目交付物','mode':'llm'})
        self.assertEqual(result.status_code,503)
        self.assertIn('未就绪',result.json()['detail'])
        self.assertEqual(self.client.get(self.path+'/runs').json()[0]['output']['status'],503)

    def test_model_citation_validation_with_test_double(self):
        self.upload()
        with server.db() as conn:source=server.source_chunks(conn,self.project)[0]
        valid={'abstained':False,'claims':[{'text':'交付现状问题地图。','source_id':source['id'],'quote':source['text']}]}
        # Test double checks software contracts ONLY, never reported as a real model evaluation.
        with patch('app.model_status',return_value={'ready':True}),patch('app.generate',return_value=(valid,{'model':'TEST_DOUBLE'})):
            r=self.client.post(self.path+'/ask',json={'question':'项目交付物','mode':'llm'}).json()
            self.assertFalse(r['abstained'])
        for invalid in [{'abstained':False,'claims':[{'text':'虚构','source_id':source['id'],'quote':'不存在的原文'}]},
                        {'abstained':False,'claims':[{'text':'虚构','source_id':99999,'quote':source['text']}]},
                        {'abstained':False,'claims':[]},{'abstained':False,'claims':['bad']},
                        {'abstained':True,'claims':[]}]:
            with patch('app.model_status',return_value={'ready':True}),patch('app.generate',return_value=(invalid,{})):
                self.assertTrue(self.client.post(self.path+'/ask',json={'question':'项目交付物','mode':'llm'}).json()['abstained'])

    def test_model_task_validation_with_test_double(self):
        invalid={'items':[{'title':'采购无人机','owner':'虚构人','due_date':None,'quote':'会议讨论预算'}]}
        with patch('app.generate',return_value=(invalid,{})):
            self.assertEqual(self.client.post(self.path+'/drafts',json={'text':'会议讨论预算','mode':'llm'}).status_code,422)
        self.assertEqual(self.client.get(self.path+'/tasks').json(),[])

    def test_model_endpoint_configuration(self):
        with patch.dict(os.environ,{'LLM_ENABLED':'1','LLM_MODEL':'test','LLM_BASE_URL':'http://remote.example/v1','LLM_API_KEY':'not-a-real-key'}):
            self.assertFalse(model_status()['ready'])
        with patch.dict(os.environ,{'LLM_ENABLED':'1','LLM_MODEL':'test','LLM_BASE_URL':'http://127.0.0.1:11434/v1','LLM_API_KEY':''}):
            self.assertTrue(model_status()['ready'])

    def test_persistence_across_connections(self):
        doc=self.upload().json()
        with TestClient(server.app,headers={'X-Assistant-Request':'1'}) as second:
            self.assertEqual(second.get(self.path+'/documents').json()[0]['id'],doc['id'])

    def test_provider_http_contract_without_network(self):
        real_client = httpx.Client
        def reply(request):
            import json
            self.assertEqual(str(request.url),'http://127.0.0.1:11434/v1/chat/completions')
            payload=json.loads(request.content)
            self.assertEqual(payload['response_format'],{'type':'json_object'})
            self.assertEqual(payload['messages'][0]['role'],'system')
            return httpx.Response(200,json={'choices':[{'message':{'content':'{"abstained":true,"claims":[]}'}}],'usage':{'total_tokens':12}})
        with patch.dict(os.environ,{'LLM_ENABLED':'1','LLM_MODEL':'test-contract','LLM_BASE_URL':'http://127.0.0.1:11434/v1','LLM_API_KEY':''}),patch('model.httpx.Client',side_effect=lambda **kw:real_client(transport=httpx.MockTransport(reply),**kw)):
            result,meta=generate('test prompt',{'question':'test'})
            self.assertTrue(result['abstained']);self.assertEqual(meta['usage']['total_tokens'],12)
        with patch('model.model_status',return_value={'ready':True,'model':'test'}),patch.dict(os.environ,{'LLM_BASE_URL':'http://127.0.0.1:11434/v1'}),patch('model.httpx.Client',side_effect=lambda **kw:real_client(transport=httpx.MockTransport(lambda r:httpx.Response(429,text='SECRET_PROVIDER_ERROR')),**kw)):
            with self.assertRaises(server.HTTPException) as caught:generate('test',{})
            self.assertEqual(caught.exception.status_code,502)
            self.assertNotIn('SECRET_PROVIDER_ERROR',caught.exception.detail)


if __name__=='__main__':
    unittest.main(verbosity=2)
